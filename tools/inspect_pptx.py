# -*- coding: utf-8 -*-
"""Dump the structure of ds-capstone-template.pptx: slides, shapes, text, picture slots."""
import io
import sys

from pptx import Presentation
from pptx.util import Emu

PATH = r"C:\Users\khali\Downloads\Claud_IBM-capstone-story\ds-capstone-template.pptx"
OUT = r"C:\Users\khali\Downloads\Claud_IBM-capstone-story\pptx_structure.txt"
sys.stdout = open(OUT, "w", encoding="utf-8")
prs = Presentation(PATH)
print("slides:", len(prs.slides))
print("size (EMU):", prs.slide_width, prs.slide_height,
      "=> inches:", round(prs.slide_width / 914400, 2), "x", round(prs.slide_height / 914400, 2))
print()

for i, slide in enumerate(prs.slides):
    print(f"===== SLIDE {i+1} (layout: {slide.slide_layout.name}) =====")
    for sh in slide.shapes:
        kind = sh.shape_type
        name = sh.name
        try:
            l, t, w, h = sh.left, sh.top, sh.width, sh.height
            geo = f"L{Emu(l).inches:.2f} T{Emu(t).inches:.2f} W{Emu(w).inches:.2f} H{Emu(h).inches:.2f}"
        except Exception:
            geo = "?"
        txt = ""
        if sh.has_text_frame:
            txt = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
        ph = ""
        if sh.is_placeholder:
            ph = f" PLACEHOLDER idx={sh.placeholder_format.idx} type={sh.placeholder_format.type}"
        print(f"  - [{name}] {kind} {geo}{ph}")
        if txt:
            print(f"      TEXT: {txt[:400]}")
    print()
